import io
import re
import logging
from typing import Dict, Any, Optional

# ReportLab for PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)

# python-docx for DOCX generation
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.resume_template_service import ResumeTemplateService

logger = logging.getLogger(__name__)

def hex_to_reportlab_color(hex_str: str, default="#1e3a8a"):
    try:
        clean = (hex_str or default).lstrip("#")
        if len(clean) == 6:
            r = int(clean[0:2], 16) / 255.0
            g = int(clean[2:4], 16) / 255.0
            b = int(clean[4:6], 16) / 255.0
            return colors.Color(r, g, b)
    except Exception:
        pass
    return colors.HexColor(default)

def hex_to_rgb(hex_str: str, default="1E3A8A"):
    try:
        clean = (hex_str or default).lstrip("#")
        if len(clean) == 6:
            return RGBColor(int(clean[0:2], 16), int(clean[2:4], 16), int(clean[4:6], 16))
    except Exception:
        pass
    return RGBColor(30, 58, 138)

class ResumeExportService:
    """
    High-fidelity PDF and DOCX resume generation engine.
    Produces selectable, structured, ATS-compliant documents matching preview formatting.
    """

    @staticmethod
    def _clean_str(val: Any, default: str = "") -> str:
        """Sanitizes any string, dict, or stringified dict into a clean display string."""
        if val is None:
            return default
        if isinstance(val, dict):
            if "value" in val and val["value"] is not None:
                return ResumeExportService._clean_str(val["value"], default)
            if "name" in val and val["name"] is not None:
                return ResumeExportService._clean_str(val["name"], default)
            return default
        if isinstance(val, (int, float)):
            return str(val)
        s = str(val).strip()
        if "{" in s and ("value" in s or "source" in s):
            s = re.sub(r"\{['\"]value['\"]:\s*['\"]?([^,'\"}]+)['\"]?[^}]*\}", r"\1", s)
        s = re.sub(r'(?i)\bCGPA\s+CGPA\b', 'CGPA', s)
        return s.strip()

    @classmethod
    def generate_pdf(cls, resume_data: dict, template_id: str = "ats_classic", formatting: Optional[dict] = None) -> io.BytesIO:
        """
        Builds a professional, ATS-friendly PDF document using ReportLab Platypus.
        """
        formatting = formatting or {}
        accent_hex = formatting.get("accent_color") or "#0f172a"
        if template_id == "ats_classic":
            accent_hex = "#0f172a"  # Classic single-color ATS compliance
        
        accent_color = hex_to_reportlab_color(accent_hex, default="#0f172a")

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Base typography adjustments
        font_size_mode = formatting.get("font_size", "standard")
        scale = 1.0
        if font_size_mode == "compact":
            scale = 0.92
        elif font_size_mode == "relaxed":
            scale = 1.08

        name_style = ParagraphStyle(
            "CandidateName",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=20 * scale,
            leading=24 * scale,
            textColor=accent_color,
            alignment=0  # Left-aligned
        )

        contact_style = ParagraphStyle(
            "ContactLine",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9 * scale,
            leading=13 * scale,
            textColor=colors.HexColor("#334155")
        )

        section_heading_style = ParagraphStyle(
            "SectionHeading",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11 * scale,
            leading=15 * scale,
            textColor=accent_color,
            spaceAfter=4,
            spaceBefore=10
        )

        body_style = ParagraphStyle(
            "BodyText",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5 * scale,
            leading=13.5 * scale,
            textColor=colors.HexColor("#1e293b")
        )

        subhead_bold = ParagraphStyle(
            "SubheadBold",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10 * scale,
            leading=14 * scale,
            textColor=colors.HexColor("#0f172a")
        )

        subhead_meta = ParagraphStyle(
            "SubheadMeta",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=9 * scale,
            leading=13 * scale,
            textColor=colors.HexColor("#475569")
        )

        bullet_style = ParagraphStyle(
            "BulletText",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9 * scale,
            leading=13 * scale,
            leftIndent=14,
            firstLineIndent=-10,
            textColor=colors.HexColor("#1e293b")
        )

        story = []

        # 1. Header (Name & Contact)
        personal = resume_data.get("personal", {})
        raw_name = personal.get("fullName") or personal.get("full_name") or "Your Name"
        full_name = cls._clean_str(raw_name)
        story.append(Paragraph(full_name, name_style))

        contact_parts = []
        if personal.get("email"):
            contact_parts.append(cls._clean_str(personal["email"]))
        if personal.get("phone"):
            contact_parts.append(cls._clean_str(personal["phone"]))
        if personal.get("location"):
            contact_parts.append(cls._clean_str(personal["location"]))
        if personal.get("linkedin"):
            clean_li = cls._clean_str(personal["linkedin"]).replace("https://", "").replace("http://", "")
            contact_parts.append(f'<a href="{personal["linkedin"]}" color="#2563eb">{clean_li}</a>')
        if personal.get("github"):
            clean_gh = cls._clean_str(personal["github"]).replace("https://", "").replace("http://", "")
            contact_parts.append(f'<a href="{personal["github"]}" color="#2563eb">{clean_gh}</a>')
        if personal.get("portfolio"):
            clean_pf = cls._clean_str(personal["portfolio"]).replace("https://", "").replace("http://", "")
            contact_parts.append(f'<a href="{personal["portfolio"]}" color="#2563eb">{clean_pf}</a>')

        contact_text = " • ".join(contact_parts)
        if contact_text:
            story.append(Spacer(1, 3))
            story.append(Paragraph(contact_text, contact_style))

        story.append(Spacer(1, 6))
        story.append(HRFlowable(width="100%", thickness=1.2, color=accent_color, spaceBefore=4, spaceAfter=8))

        # Check visible sections
        vis = formatting.get("visible_sections", {})

        # Order sections according to user customization or template default
        template_info = ResumeTemplateService.get_template(template_id)
        section_order = formatting.get("section_order") or template_info.get("section_order", [
            "summary", "skills", "projects", "experience", "education", "certifications", "achievements"
        ])

        for section_key in section_order:
            if vis.get(section_key) is False:
                continue

            # Section: Summary
            if section_key == "summary" and resume_data.get("summary"):
                story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading_style))
                story.append(Paragraph(cls._clean_str(resume_data["summary"]), body_style))
                story.append(Spacer(1, 6))

            # Section: Education
            elif section_key == "education" and resume_data.get("education"):
                story.append(Paragraph("EDUCATION", section_heading_style))
                for edu in resume_data["education"]:
                    inst = cls._clean_str(edu.get("institution") or edu.get("college") or "")
                    deg = cls._clean_str(edu.get("degree") or "")
                    yr = cls._clean_str(edu.get("year") or edu.get("graduationYear") or "")
                    gpa_raw = cls._clean_str(edu.get("gpa") or edu.get("cgpa") or "")
                    gpa = gpa_raw
                    if gpa and "cgpa" not in gpa.lower() and "gpa" not in gpa.lower():
                        gpa = f"{gpa} CGPA"

                    header_line = f"<b>{deg}</b>"
                    if inst:
                        header_line += f" — {inst}"
                    meta_line = " | ".join([m for m in [yr, gpa] if m])

                    row = [
                        [Paragraph(header_line, subhead_bold), Paragraph(meta_line, ParagraphStyle("RAlign", parent=subhead_meta, alignment=2))]
                    ]
                    t = Table(row, colWidths=[410, 130])
                    t.setStyle(TableStyle([
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                        ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ]))
                    story.append(t)
                story.append(Spacer(1, 6))

            # Section: Skills
            elif section_key == "skills" and resume_data.get("skills"):
                story.append(Paragraph("TECHNICAL SKILLS", section_heading_style))
                skills = resume_data["skills"]
                if isinstance(skills, dict):
                    for cat_name, skill_list in skills.items():
                        if skill_list:
                            clean_skills = [cls._clean_str(s) for s in skill_list if cls._clean_str(s)]
                            line = f"<b>{cat_name}:</b> {', '.join(clean_skills)}"
                elif isinstance(skills, list):
                    clean_skills = [cls._clean_str(s) for s in skills if cls._clean_str(s)]
                    line = f"<b>Core Competencies:</b> {', '.join(clean_skills)}"
                    story.append(Paragraph(line, body_style))
                story.append(Spacer(1, 6))

            # Section: Projects
            elif section_key == "projects" and resume_data.get("projects"):
                story.append(Paragraph("PROJECTS", section_heading_style))
                for proj in resume_data["projects"]:
                    name = proj.get("name") or "Project"
                    role = proj.get("role") or ""
                    techs = proj.get("technologies") or []
                    if isinstance(techs, str):
                        techs = [t.strip() for t in techs.split(",")]
                    tech_str = f"({', '.join(techs)})" if techs else ""

                    proj_header = f"<b>{name}</b> {tech_str}".strip()
                    story.append(Paragraph(proj_header, subhead_bold))

                    bullets = proj.get("bullets", [])
                    if not bullets and proj.get("description"):
                        bullets = [proj["description"]]

                    for b in bullets:
                        if b and str(b).strip():
                            story.append(Paragraph(f"• {str(b).strip()}", bullet_style))
                    story.append(Spacer(1, 4))
                story.append(Spacer(1, 4))

            # Section: Experience
            elif section_key == "experience" and resume_data.get("experience"):
                story.append(Paragraph("WORK & LEADERSHIP EXPERIENCE", section_heading_style))
                for exp in resume_data["experience"]:
                    role = exp.get("role") or "Role"
                    company = exp.get("company") or exp.get("organization") or ""
                    duration = exp.get("duration") or ""
                    if not duration:
                        dates = [exp.get("startDate"), exp.get("endDate")]
                        duration = " - ".join([d for d in dates if d])

                    header_line = f"<b>{role}</b>"
                    if company:
                        header_line += f" — {company}"

                    row = [
                        [Paragraph(header_line, subhead_bold), Paragraph(duration, ParagraphStyle("RAlignExp", parent=subhead_meta, alignment=2))]
                    ]
                    t = Table(row, colWidths=[410, 130])
                    t.setStyle(TableStyle([
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                        ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ]))
                    story.append(t)

                    bullets = exp.get("bullets", [])
                    if not bullets and exp.get("responsibilities"):
                        bullets = exp["responsibilities"]
                    for b in bullets:
                        if b and str(b).strip():
                            story.append(Paragraph(f"• {str(b).strip()}", bullet_style))
                    story.append(Spacer(1, 4))
                story.append(Spacer(1, 4))

            # Section: Certifications
            elif section_key == "certifications" and resume_data.get("certifications"):
                story.append(Paragraph("CERTIFICATIONS", section_heading_style))
                for cert in resume_data["certifications"]:
                    name = cert.get("name") if isinstance(cert, dict) else str(cert)
                    org = cert.get("organization") if isinstance(cert, dict) else ""
                    date = cert.get("date") if isinstance(cert, dict) else ""
                    line = f"• <b>{name}</b>"
                    if org:
                        line += f" — {org}"
                    if date:
                        line += f" ({date})"
                    story.append(Paragraph(line, body_style))
                story.append(Spacer(1, 6))

            # Section: Achievements
            elif section_key == "achievements" and resume_data.get("achievements"):
                story.append(Paragraph("HONORS & ACHIEVEMENTS", section_heading_style))
                for ach in resume_data["achievements"]:
                    title = ach.get("title") if isinstance(ach, dict) else str(ach)
                    ctx = ach.get("context") if isinstance(ach, dict) else ""
                    line = f"• <b>{title}</b>"
                    if ctx:
                        line += f": {ctx}"
                    story.append(Paragraph(line, body_style))
                story.append(Spacer(1, 6))

        doc.build(story)
        buffer.seek(0)
        return buffer

    @classmethod
    def generate_docx(cls, resume_data: dict, template_id: str = "ats_classic", formatting: Optional[dict] = None) -> io.BytesIO:
        """
        Builds a standard, editable ATS-friendly DOCX document using python-docx.
        """
        formatting = formatting or {}
        accent_hex = formatting.get("accent_color") or "#0F172A"
        accent_rgb = hex_to_rgb(accent_hex, default="0F172A")

        doc = docx.Document()

        # Set 0.5 in margins for compact professional presentation
        for s in doc.sections:
            s.top_margin = Inches(0.5)
            s.bottom_margin = Inches(0.5)
            s.left_margin = Inches(0.5)
            s.right_margin = Inches(0.5)

        # 1. Candidate Name
        personal = resume_data.get("personal", {})
        full_name = personal.get("fullName") or personal.get("full_name") or "Your Name"

        name_p = doc.add_paragraph()
        name_p.paragraph_format.space_before = Pt(0)
        name_p.paragraph_format.space_after = Pt(2)
        run_name = name_p.add_run(full_name)
        run_name.font.size = Pt(20)
        run_name.font.bold = True
        run_name.font.color.rgb = accent_rgb

        # Contact info
        contact_parts = []
        if personal.get("email"):
            contact_parts.append(personal["email"])
        if personal.get("phone"):
            contact_parts.append(personal["phone"])
        if personal.get("location"):
            contact_parts.append(personal["location"])
        if personal.get("linkedin"):
            contact_parts.append(personal["linkedin"])
        if personal.get("github"):
            contact_parts.append(personal["github"])

        if contact_parts:
            contact_p = doc.add_paragraph(" | ".join(contact_parts))
            contact_p.paragraph_format.space_before = Pt(0)
            contact_p.paragraph_format.space_after = Pt(8)
            contact_p.runs[0].font.size = Pt(9.5)
            contact_p.runs[0].font.color.rgb = RGBColor(71, 85, 105)

        def add_heading(title_text):
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(10)
            hp.paragraph_format.space_after = Pt(3)
            r = hp.add_run(title_text.upper())
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = accent_rgb

        # Sections according to user customization or template default
        template_info = ResumeTemplateService.get_template(template_id)
        section_order = formatting.get("section_order") or template_info.get("section_order", [
            "summary", "skills", "projects", "experience", "education", "certifications", "achievements"
        ])

        vis = formatting.get("visible_sections", {})

        for section_key in section_order:
            if vis.get(section_key) is False:
                continue

            # Summary
            if section_key == "summary" and resume_data.get("summary"):
                add_heading("Professional Summary")
                sp = doc.add_paragraph(cls._clean_str(resume_data["summary"]))
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)
                sp.runs[0].font.size = Pt(10)

            # Education
            elif section_key == "education" and resume_data.get("education"):
                add_heading("Education")
                for edu in resume_data["education"]:
                    deg = cls._clean_str(edu.get("degree") or "")
                    inst = cls._clean_str(edu.get("institution") or edu.get("college") or "")
                    yr = cls._clean_str(edu.get("year") or edu.get("graduationYear") or "")
                    gpa_raw = cls._clean_str(edu.get("gpa") or edu.get("cgpa") or "")
                    gpa = gpa_raw
                    if gpa and "cgpa" not in gpa.lower() and "gpa" not in gpa.lower():
                        gpa = f"{gpa} CGPA"

                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    r1 = p.add_run(f"{deg} — {inst}" if inst else deg)
                    r1.font.bold = True
                    r1.font.size = Pt(10)

                    meta = " | ".join([m for m in [yr, gpa] if m])
                    if meta:
                        r2 = p.add_run(f" ({meta})")
                        r2.font.italic = True
                        r2.font.size = Pt(9.5)

            # Skills
            elif section_key == "skills" and resume_data.get("skills"):
                add_heading("Technical Skills")
                skills = resume_data["skills"]
                if isinstance(skills, dict):
                    for cat, s_list in skills.items():
                        if s_list:
                            clean_s = [cls._clean_str(s) for s in s_list if cls._clean_str(s)]
                            sp = doc.add_paragraph()
                            sp.paragraph_format.space_before = Pt(1)
                            sp.paragraph_format.space_after = Pt(1)
                            r_cat = sp.add_run(f"{cat}: ")
                            r_cat.font.bold = True
                            r_cat.font.size = Pt(9.5)
                            r_val = sp.add_run(", ".join(clean_s))
                            r_val.font.size = Pt(9.5)
                elif isinstance(skills, list):
                    clean_s = [cls._clean_str(s) for s in skills if cls._clean_str(s)]
                    sp = doc.add_paragraph(", ".join(clean_s))
                    sp.paragraph_format.space_before = Pt(1)
                    sp.paragraph_format.space_after = Pt(1)
                    sp.runs[0].font.size = Pt(9.5)

            # Projects
            elif section_key == "projects" and resume_data.get("projects"):
                add_heading("Projects")
                for proj in resume_data["projects"]:
                    name = proj.get("name") or "Project"
                    techs = proj.get("technologies") or []
                    if isinstance(techs, str):
                        techs = [t.strip() for t in techs.split(",")]
                    tech_str = f" ({', '.join(techs)})" if techs else ""

                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(1)
                    r_p = p.add_run(f"{name}{tech_str}")
                    r_p.font.bold = True
                    r_p.font.size = Pt(10)

                    bullets = proj.get("bullets", [])
                    if not bullets and proj.get("description"):
                        bullets = [proj["description"]]
                    for b in bullets:
                        if b and str(b).strip():
                            bp = doc.add_paragraph(str(b).strip(), style="List Bullet")
                            bp.paragraph_format.space_before = Pt(0)
                            bp.paragraph_format.space_after = Pt(1)
                            bp.runs[0].font.size = Pt(9.5)

            # Experience
            elif section_key == "experience" and resume_data.get("experience"):
                add_heading("Experience")
                for exp in resume_data["experience"]:
                    role = exp.get("role") or "Role"
                    comp = exp.get("company") or exp.get("organization") or ""
                    dur = exp.get("duration") or ""
                    if not dur:
                        dates = [exp.get("startDate"), exp.get("endDate")]
                        dur = " - ".join([d for d in dates if d])

                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(1)
                    r_e = p.add_run(f"{role} — {comp}")
                    r_e.font.bold = True
                    r_e.font.size = Pt(10)
                    if dur:
                        r_d = p.add_run(f" ({dur})")
                        r_d.font.italic = True
                        r_d.font.size = Pt(9)

                    bullets = exp.get("bullets", [])
                    if not bullets and exp.get("responsibilities"):
                        bullets = exp["responsibilities"]
                    for b in bullets:
                        if b and str(b).strip():
                            bp = doc.add_paragraph(str(b).strip(), style="List Bullet")
                            bp.paragraph_format.space_before = Pt(0)
                            bp.paragraph_format.space_after = Pt(1)
                            bp.runs[0].font.size = Pt(9.5)

            # Certifications
            elif section_key == "certifications" and resume_data.get("certifications"):
                add_heading("Certifications")
                for cert in resume_data["certifications"]:
                    name = cert.get("name") if isinstance(cert, dict) else str(cert)
                    org = cert.get("organization") if isinstance(cert, dict) else ""
                    text = f"{name} ({org})" if org else name
                    bp = doc.add_paragraph(text, style="List Bullet")
                    bp.paragraph_format.space_before = Pt(0)
                    bp.paragraph_format.space_after = Pt(1)
                    bp.runs[0].font.size = Pt(9.5)

            # Achievements
            elif section_key == "achievements" and resume_data.get("achievements"):
                add_heading("Achievements")
                for ach in resume_data["achievements"]:
                    title = ach.get("title") if isinstance(ach, dict) else str(ach)
                    ctx = ach.get("context") if isinstance(ach, dict) else ""
                    text = f"{title}: {ctx}" if ctx else title
                    bp = doc.add_paragraph(text, style="List Bullet")
                    bp.paragraph_format.space_before = Pt(0)
                    bp.paragraph_format.space_after = Pt(1)
                    bp.runs[0].font.size = Pt(9.5)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer
